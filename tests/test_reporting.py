"""Report integration checks: one active batch, two files, unchanged KPI facts."""

from __future__ import annotations

from dataclasses import replace
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook
from reportlab.platypus.doctemplate import LayoutError
from streamlit.testing.v1 import AppTest

from manufacturing_operations_intelligence.domain.analytics import KpiScope
from manufacturing_operations_intelligence.reporting import render_excel, render_pdf
from manufacturing_operations_intelligence.services.application import (
    ApplicationServiceError,
    ManufacturingApplicationService,
)
from manufacturing_operations_intelligence.summaries import generate_management_summary

SCOPE = KpiScope("2025-01-01", "2025-03-31", line_ids=("LINE-02",))
APP_FILE = Path(__file__).resolve().parents[1] / "streamlit_app.py"


@pytest.fixture
def service(tmp_path: Path) -> ManufacturingApplicationService:
    instance = ManufacturingApplicationService.from_path(
        tmp_path / "report-test.sqlite3", downtime_above_minutes=Decimal("120")
    )
    assert instance.load_sample_data().status == "accepted"
    return instance


def test_management_files_and_sections_match_analytics(service, tmp_path: Path) -> None:
    engine = service.get_overview_metrics(SCOPE)
    payload = service.get_report_payload(SCOPE)
    artifacts = service.export_management_reports(SCOPE)
    excel_path = tmp_path / "management.xlsx"
    pdf_path = tmp_path / "management.pdf"
    excel_path.write_bytes(artifacts.excel)
    pdf_path.write_bytes(artifacts.pdf)

    assert artifacts.identity.batch_id == engine.identity.batch_id == payload.identity.batch_id
    assert artifacts.scope == SCOPE
    assert excel_path.stat().st_size > 5_000
    assert pdf_path.stat().st_size > 5_000
    assert artifacts.excel.startswith(b"PK")
    assert artifacts.pdf.startswith(b"%PDF")

    workbook = load_workbook(excel_path, read_only=False, data_only=False)
    assert workbook.sheetnames == [
        "Overview",
        "Production",
        "Quality",
        "Inventory",
        "Anomalies",
        "Data Quality",
        "Summary",
    ]
    overview = {
        row[0]: row[1] for row in workbook["Overview"].iter_rows(min_row=3, values_only=True)
    }
    for label, key in (
        ("Production attainment", "KPI-PA"),
        ("Good yield", "KPI-GY"),
        ("Scrap rate", "KPI-SR"),
        ("Total downtime", "KPI-TD"),
        ("Completed orders", "KPI-CO"),
        ("Schedule adherence", "KPI-SA"),
        ("Materials below safety stock", "KPI-IR"),
    ):
        expected = engine.metrics[key]
        assert overview[label] == (expected.display_value if expected.status == "valid" else "N/A")
        if expected.status == "valid":
            assert expected.display_value.encode() in artifacts.pdf
    assert overview["Report period"] == "2025-01-01 to 2025-03-31 (inclusive)"
    assert overview["Production lines"] == "LINE-02"
    assert "line and product filters do not apply" in overview["Inventory scope"]
    assert len(workbook["Production"]._charts) == 1
    assert len(workbook["Quality"]._charts) == 1
    assert workbook["Production"]["B3"].value == payload.series.daily[0].planned_qty
    assert workbook["Quality"]["C3"].value == payload.series.daily[0].scrap_qty
    assert workbook["Data Quality"]["B7"].value == "unsupported_schema"

    for section in (
        b"Management report",
        b"Executive KPI summary",
        b"Production analysis",
        b"Quality analysis",
        b"Inventory risks",
        b"Detected anomalies",
        b"Data quality and coverage",
        b"Management summary",
    ):
        assert section in artifacts.pdf
    assert artifacts.pdf.count(b"/Type /Page") >= 2


def test_anomaly_detail_is_complete_in_excel_and_pdf_preview_is_labeled(service) -> None:
    scope = KpiScope("2025-01-01", "2025-03-31")
    payload = service.get_report_payload(scope)
    assert payload.detection is not None
    assert payload.detection.anomaly_count > 20
    artifacts = service.export_management_reports(scope)
    workbook = load_workbook(BytesIO(artifacts.excel), read_only=True)
    sheet = workbook["Anomalies"]
    assert sheet.max_row - 2 == payload.detection.anomaly_count
    assert workbook["Inventory"].max_row - 2 == len(payload.metrics["KPI-IR"].details)
    assert f"Showing 20 of {payload.detection.anomaly_count} anomalies".encode() in artifacts.pdf
    assert b"Page 1" in artifacts.pdf
    assert payload.identity.batch_id[:12].encode() in artifacts.pdf
    assert b"20550/659" not in artifacts.pdf  # Exact fractions are formatted for managers.
    assert "/" not in str(sheet["F3"].value)


def test_current_management_summary_is_included_and_stale_summary_is_rejected(service) -> None:
    payload = service.get_report_payload(SCOPE)
    summary = generate_management_summary(
        payload,
        enabled=True,
        api_key="test-key",
        model="test-model",
        provider=lambda *_args: '{"observations":["KPI-PA"],"management_attention":["KPI-SR"]}',
    )
    assert summary.status == "ai_draft"

    artifacts = service.export_management_reports(SCOPE, summary=summary)
    workbook = load_workbook(BytesIO(artifacts.excel), read_only=True)
    summary_rows = list(workbook["Summary"].iter_rows(min_row=3, values_only=True))
    assert ("Summary type", "AI-assisted draft — human review required") in summary_rows
    assert any(summary.observations[0] == row[1] for row in summary_rows)
    assert summary.observations[0].encode() in artifacts.pdf
    assert b"KPI definitions" in artifacts.pdf
    assert b"Gross actual output divided by planned output" in artifacts.pdf

    stale = replace(summary, scope=KpiScope("2025-01-02", "2025-03-31"))
    with pytest.raises(ApplicationServiceError, match="stale"):
        service.export_management_reports(SCOPE, summary=stale)


def test_unconfigured_rules_and_empty_production_are_labeled(service, tmp_path: Path) -> None:
    no_rules = ManufacturingApplicationService.from_path(tmp_path / "report-test.sqlite3")
    scope = KpiScope("2026-01-01", "2026-01-02")
    artifacts = no_rules.export_management_reports(scope)
    workbook = load_workbook(BytesIO(artifacts.excel), read_only=True)
    overview = {
        row[0]: row[1] for row in workbook["Overview"].iter_rows(min_row=3, values_only=True)
    }
    assert overview["Production attainment"] == "N/A"
    assert overview["Anomaly checks"].startswith("Not evaluated")
    assert workbook["Production"]["A3"].value == "No selected production records"
    assert b"No selected records" in artifacts.pdf
    assert b"not evaluated" in artifacts.pdf.lower()


def test_renderers_use_supplied_values_and_escape_source_formulas(service) -> None:
    payload = service.get_report_payload(SCOPE)
    attainment = replace(payload.metrics["KPI-PA"], display_value="66.66%")
    inventory = replace(
        payload.metrics["KPI-IR"],
        details=({**payload.metrics["KPI-IR"].details[0], "material_id": '=HYPERLINK("x")'},),
    )
    modified = replace(
        payload, metrics={**payload.metrics, "KPI-PA": attainment, "KPI-IR": inventory}
    )
    excel = render_excel(modified)
    pdf = render_pdf(modified)
    workbook = load_workbook(BytesIO(excel), read_only=False)
    overview = {
        row[0]: row[1] for row in workbook["Overview"].iter_rows(min_row=3, values_only=True)
    }
    assert overview["Production attainment"] == "66.66%"
    assert b"66.66%" in pdf
    assert workbook["Inventory"]["A3"].value == '=HYPERLINK("x")'
    assert workbook["Inventory"]["A3"].data_type == "s"


def test_no_active_batch_and_stale_payload_are_rejected(tmp_path: Path, service) -> None:
    empty = ManufacturingApplicationService.from_path(tmp_path / "empty.sqlite3")
    with pytest.raises(ApplicationServiceError, match="Load a validated dataset"):
        empty.export_management_reports(SCOPE)
    payload = service.get_report_payload(SCOPE)
    stale_metric = replace(payload.metrics["KPI-PA"], batch_id="another-batch")
    stale = replace(payload, metrics={**payload.metrics, "KPI-PA": stale_metric})
    with pytest.raises(ValueError, match="does not match"):
        render_excel(stale)
    with pytest.raises(ValueError, match="does not match"):
        render_pdf(stale)

    stale_fingerprint = replace(
        payload,
        metrics={
            **payload.metrics,
            "KPI-PA": replace(payload.metrics["KPI-PA"], fingerprint="stale-fingerprint"),
        },
    )
    with pytest.raises(ValueError, match="identity"):
        render_excel(stale_fingerprint)
    with pytest.raises(ValueError, match="identity"):
        render_pdf(stale_fingerprint)


def test_many_selected_lines_fit_paginated_pdf(service) -> None:
    payload = service.get_report_payload(SCOPE)
    scope = replace(SCOPE, line_ids=tuple(f"LINE-{index:04d}" for index in range(500)))
    metrics = {key: replace(value, scope=scope) for key, value in payload.metrics.items()}
    detection = replace(payload.detection, scope=scope)
    scoped = replace(payload, scope=scope, metrics=metrics, detection=detection)
    pdf = render_pdf(scoped)
    assert pdf.startswith(b"%PDF")
    assert b"LINE-0499" in pdf
    assert pdf.count(b"/Type /Page") < 10


def test_pdf_failure_preserves_generated_excel(service, monkeypatch) -> None:
    def fail_pdf(_payload):
        raise LayoutError("oversized scope")

    monkeypatch.setattr(
        "manufacturing_operations_intelligence.services.application.render_pdf", fail_pdf
    )
    artifacts = service.export_management_reports(SCOPE)
    assert artifacts.excel is not None and artifacts.excel.startswith(b"PK")
    assert artifacts.pdf is None
    assert artifacts.pdf_error == "PDF generation failed for this selection."


def test_excel_failure_preserves_generated_pdf(service, monkeypatch) -> None:
    def fail_excel(_payload):
        raise OSError("workbook writer failed")

    monkeypatch.setattr(
        "manufacturing_operations_intelligence.services.application.render_excel", fail_excel
    )
    artifacts = service.export_management_reports(SCOPE)
    assert artifacts.pdf is not None and artifacts.pdf.startswith(b"%PDF")
    assert artifacts.excel is None
    assert artifacts.excel_error == "Excel generation failed for this selection."


def test_dashboard_downloads_are_prepared_on_demand_and_expire_with_filters(
    service, tmp_path: Path, monkeypatch
) -> None:
    monkeypatch.setenv("MOI_DATABASE_PATH", str(tmp_path / "report-test.sqlite3"))
    app = AppTest.from_file(APP_FILE, default_timeout=10).run()
    next(widget for widget in app.radio if widget.label == "Area").set_value(
        "Reports & Insights"
    ).run()
    assert not app.exception
    assert "prepared_report" not in app.session_state
    next(
        widget for widget in app.button if widget.label == "Generate management reports"
    ).click().run()
    assert not app.exception
    artifacts = app.session_state["prepared_report"][1]
    assert artifacts.excel.startswith(b"PK") and artifacts.pdf.startswith(b"%PDF")
    app.session_state["prepared_report"] = (("old-definition",), artifacts)
    app.run()
    assert "prepared_report" not in app.session_state
    next(
        widget for widget in app.button if widget.label == "Generate management reports"
    ).click().run()
    assert "prepared_report" in app.session_state
    next(widget for widget in app.date_input if widget.label == "From").set_value(
        "2025-02-01"
    ).run()
    assert not app.exception
    assert "prepared_report" not in app.session_state
