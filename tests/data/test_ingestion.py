"""Input contract checks using small synthetic four-domain attempts."""

from __future__ import annotations

import csv
import io
import re
import zipfile
from copy import deepcopy
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import Workbook

from manufacturing_operations_intelligence.data import readers
from manufacturing_operations_intelligence.domain.validation import SCHEMAS
from manufacturing_operations_intelligence.services.ingestion import (
    ingest_csv_batch,
    ingest_excel_workbook,
)


def rows() -> dict[str, list[dict[str, object]]]:
    common = {
        "production_date": "2025-01-01",
        "line_id": "LINE-01",
        "shift_id": "DAY",
        "order_id": "MO-001",
        "product_id": "PRD-001",
        "qty_unit": "ea",
    }
    return {
        "production_plan": [
            {**common, "planned_qty": "100", "planned_production_minutes": "480.000"}
        ],
        "production_actual": [
            {
                **common,
                "actual_qty": "90",
                "runtime_minutes": "450.000",
                "downtime_minutes": "20.000",
            }
        ],
        "quality": [{**common, "good_qty": "88", "scrap_qty": "2"}],
        "inventory": [
            {
                "snapshot_date": "2025-01-01",
                "material_id": "MAT-001",
                "qty_unit": "kg",
                "inventory_qty": "120.500",
                "safety_stock": "100.000",
            }
        ],
    }


def csv_files(
    data: dict[str, list[dict[str, object]]],
    *,
    headers: dict[str, list[str]] | None = None,
) -> dict[str, tuple[str, bytes]]:
    files = {}
    for dataset, records in data.items():
        output = io.StringIO(newline="")
        fieldnames = (headers or {}).get(dataset, list(SCHEMAS[dataset]))
        writer = csv.DictWriter(
            output, fieldnames=fieldnames, extrasaction="ignore", lineterminator="\n"
        )
        writer.writeheader()
        writer.writerows(records)
        files[dataset] = (f"synthetic_{dataset}.csv", output.getvalue().encode())
    return files


def workbook_bytes(data: dict[str, list[dict[str, object]]]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for dataset, records in data.items():
        sheet = workbook.create_sheet(dataset)
        sheet.append(list(SCHEMAS[dataset]))
        for record in records:
            sheet.append([record[name] for name in SCHEMAS[dataset]])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def errors(result, code: str):
    return [issue for issue in result.issues if issue.code == f"DATA.{code}"]


def test_valid_csv_and_excel_batches() -> None:
    csv_result = ingest_csv_batch(csv_files(rows()))
    excel_result = ingest_excel_workbook("synthetic.xlsx", workbook_bytes(rows()))
    for result in (csv_result, excel_result):
        assert result.accepted and result.status == "accepted"
        assert result.total_observed_rows == result.evaluated_rows == 4
        assert result.unevaluated_rows == 0 and result.error_count == 0
        assert all(len(records) == 1 for records in result.records.values())
        assert result.records["production_actual"][0].values["actual_qty"] == 90
        assert str(result.records["inventory"][0].values["inventory_qty"]) == "120.500"
    assert csv_result.records["quality"][0].source_row == 2
    assert excel_result.records["quality"][0].source_row == 2


def test_excel_typed_dates_and_numeric_quantities_normalize_losslessly() -> None:
    data = rows()
    for dataset in ("production_plan", "production_actual", "quality"):
        data[dataset][0]["production_date"] = datetime(2025, 1, 1)
    data["inventory"][0]["snapshot_date"] = datetime(2025, 1, 1)
    data["production_plan"][0]["planned_qty"] = 100.0
    data["production_actual"][0]["actual_qty"] = 90
    data["inventory"][0]["inventory_qty"] = 120.5
    result = ingest_excel_workbook("synthetic.xlsx", workbook_bytes(data))
    assert result.accepted
    assert result.records["production_plan"][0].values["production_date"] == "2025-01-01"
    assert result.records["production_plan"][0].values["planned_qty"] == 100
    assert result.records["inventory"][0].values["inventory_qty"].as_tuple().exponent == -3
    assert any(issue.normalized == "120.500" for issue in errors(result, "N06"))


def test_missing_column_is_reported_and_batch_rejected() -> None:
    fields = list(SCHEMAS["quality"])
    fields.remove("scrap_qty")
    result = ingest_csv_batch(csv_files(rows(), headers={"quality": fields}))
    assert any(
        issue.dataset == "quality" and issue.field == "scrap_qty" and issue.row is None
        for issue in errors(result, "E-SCHEMA")
    )
    assert not result.accepted and not any(result.records.values())


@pytest.mark.parametrize(
    ("dataset", "field", "value", "code"),
    [
        ("production_actual", "actual_qty", "many", "T-COUNT"),
        ("production_plan", "planned_qty", "-1", "T-COUNT"),
        ("quality", "good_qty", "N/A", "E-MISSING"),
        ("quality", "order_id", "", "E-MISSING"),
        ("inventory", "snapshot_date", "2025-02-30", "T-DATE"),
        ("production_plan", "production_date", "01/01/2025", "T-DATE"),
    ],
)
def test_field_errors_identify_dataset_row_field_and_severity(dataset, field, value, code) -> None:
    data = rows()
    data[dataset][0][field] = value
    result = ingest_csv_batch(csv_files(data))
    assert not result.accepted and not any(result.records.values())
    assert any(
        issue.dataset == dataset
        and issue.row == 2
        and issue.field == field
        and issue.severity == "ERROR"
        and issue.reason
        for issue in errors(result, code)
    )


def test_mixed_valid_invalid_rows_remain_visible_without_partial_acceptance() -> None:
    data = rows()
    for dataset in ("production_plan", "production_actual", "quality"):
        valid = deepcopy(data[dataset][0])
        valid.update(production_date="2025-01-02", order_id="MO-002")
        data[dataset].append(valid)
    data["quality"][1]["scrap_qty"] = "invalid"
    result = ingest_csv_batch(csv_files(data))
    assert not result.accepted
    assert len(result.candidates["quality"]) == 1
    assert len(result.candidates["production_plan"]) == 2
    assert not any(result.records.values())
    assert any(
        issue.dataset == "quality" and issue.row == 3 and issue.field == "scrap_qty"
        for issue in errors(result, "T-COUNT")
    )


def test_normalized_duplicate_reports_both_rows() -> None:
    data = rows()
    duplicate = deepcopy(data["inventory"][0])
    duplicate["material_id"] = " MAT-001 "
    data["inventory"].append(duplicate)
    result = ingest_csv_batch(csv_files(data))
    duplicate_issues = errors(result, "R01")
    assert {issue.row for issue in duplicate_issues} == {2, 3}
    assert {issue.counterpart_rows for issue in duplicate_issues} == {(2,), (3,)}
    assert not result.accepted


@pytest.mark.parametrize("rule", ["R02", "R03", "R04", "R05", "R06", "R07"])
def test_referential_and_consistency_rules(rule) -> None:
    data = rows()
    if rule == "R02":
        data["quality"][0]["production_date"] = "2025-01-02"
    elif rule == "R03":
        data["quality"][0]["product_id"] = "PRD-002"
    elif rule == "R04":
        data["quality"][0]["good_qty"] = "87"
    elif rule == "R05":
        data["production_actual"][0]["runtime_minutes"] = "470"
    elif rule == "R06":
        for dataset in ("production_plan", "production_actual", "quality"):
            second = deepcopy(data[dataset][0])
            second.update(shift_id="NIGHT", order_id="MO-002")
            data[dataset].append(second)
        data["production_plan"][0]["planned_production_minutes"] = "800"
        data["production_plan"][1]["planned_production_minutes"] = "800"
    else:
        second = deepcopy(data["inventory"][0])
        second.update(snapshot_date="2025-01-02", qty_unit="l")
        data["inventory"].append(second)
    result = ingest_csv_batch(csv_files(data))
    assert errors(result, rule)
    assert not result.accepted


def test_header_alias_and_literal_identifier_are_preserved() -> None:
    data = rows()
    data["production_plan"][0]["order_id"] = " 0012 "
    data["production_actual"][0]["order_id"] = " 0012 "
    data["quality"][0]["order_id"] = " 0012 "
    for dataset in ("production_plan", "production_actual", "quality"):
        data[dataset][0]["date"] = data[dataset][0].pop("production_date")
    fields = {
        dataset: ["date", *list(SCHEMAS[dataset])[1:]]
        for dataset in ("production_plan", "production_actual", "quality")
    }
    result = ingest_csv_batch(csv_files(data, headers=fields))
    assert result.accepted
    assert result.records["production_plan"][0].values["order_id"] == "0012"
    assert errors(result, "N02") and errors(result, "N03")


def test_excel_numeric_id_formula_and_non_midnight_date_are_rejected() -> None:
    data = rows()
    data["production_plan"][0]["line_id"] = 123
    payload = workbook_bytes(data)
    assert errors(ingest_excel_workbook("synthetic.xlsx", payload), "T-ID")

    data = rows()
    data["quality"][0]["scrap_qty"] = "=1+1"
    result = ingest_excel_workbook("synthetic.xlsx", workbook_bytes(data))
    assert errors(result, "E-PARSE") and not result.accepted

    data = rows()
    data["production_plan"][0]["production_date"] = datetime(2025, 1, 1, 12, 30)
    result = ingest_excel_workbook("synthetic.xlsx", workbook_bytes(data))
    assert any(issue.field == "production_date" for issue in errors(result, "T-DATE"))


def test_sparse_oversized_xlsx_dimensions_are_rejected_before_row_scan() -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for dataset in SCHEMAS:
        sheet = workbook.create_sheet(dataset)
        sheet.append(list(SCHEMAS[dataset]))
    workbook["quality"]["A100002"] = "out-of-range"
    output = io.BytesIO()
    workbook.save(output)

    result = ingest_excel_workbook("sparse.xlsx", output.getvalue())

    assert not result.accepted
    assert result.stage == "parse_incomplete"
    assert any(issue.dataset == "quality" for issue in errors(result, "E-LIMIT"))
    assert not any(result.records.values())


@pytest.mark.parametrize(
    "member",
    [
        ("xl/media/compressed.bin", b"0" * (readers.MAX_EXPANDED_BYTES + 1)),
        ("xl/worksheets/extra.xml", b"<worksheet>" + b"<c/>" *
         (readers.MAX_SHEET_CELLS + 1) + b"</worksheet>"),
    ],
)
def test_xlsx_resource_limits_precede_openpyxl_loading(monkeypatch, member) -> None:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(*member)

    def forbidden(*_args, **_kwargs):
        raise AssertionError("Unbounded workbook loader was invoked")

    monkeypatch.setattr(readers, "load_workbook", forbidden)
    _sources, problems = readers.read_workbook("compressed.xlsx", output.getvalue())
    assert problems and problems[0].code == "E-LIMIT"


def test_malformed_excel_cell_returns_validation_error() -> None:
    original = workbook_bytes(rows())
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(original)) as source, zipfile.ZipFile(output, "w") as target:
        for member in source.infolist():
            content = source.read(member.filename)
            if member.filename == "xl/worksheets/sheet3.xml":
                content, replaced = re.subn(
                    rb'<c r="A2"[^>]*>.*?</c>',
                    b'<c r="A2" t="s"><v>999999</v></c>',
                    content,
                    count=1,
                )
                assert replaced == 1
            target.writestr(member, content)

    result = ingest_excel_workbook("malformed.xlsx", output.getvalue())
    assert not result.accepted
    assert result.stage == "parse_incomplete"
    assert errors(result, "E-PARSE")


def test_oversized_csv_batch_is_rejected_before_any_parse(monkeypatch) -> None:
    def forbidden(*_args):
        raise AssertionError("An oversized batch must not be parsed")

    monkeypatch.setattr(
        "manufacturing_operations_intelligence.services.ingestion.read_csv", forbidden
    )
    content = b"0" * (readers.MAX_BYTES // 4 + 1)
    files = {dataset: (f"{dataset}.csv", content) for dataset in SCHEMAS}
    result = ingest_csv_batch(files)
    assert not result.accepted
    assert errors(result, "E-LIMIT")


def test_parser_failure_and_incomplete_batch_are_observable() -> None:
    files = csv_files(rows())
    files["inventory"] = ("synthetic_inventory.csv", b"material_id,qty_unit\n\xff")
    result = ingest_csv_batch(files)
    assert errors(result, "E-PARSE")
    assert result.total_observed_rows is None and result.stage == "parse_incomplete"
    assert not result.accepted

    result = ingest_csv_batch({"production_plan": files["production_plan"]})
    assert {issue.dataset for issue in errors(result, "E-PACKAGE")} >= {
        "production_actual",
        "quality",
        "inventory",
    }
    assert not result.accepted

    corrupt_workbook = ingest_excel_workbook("synthetic.xlsx", b"not a workbook")
    assert not corrupt_workbook.accepted
    assert corrupt_workbook.total_observed_rows is None
    assert corrupt_workbook.stage == "parse_incomplete"


def test_reserved_markers_alias_collision_and_unsafe_filename() -> None:
    data = rows()
    data["inventory"][0]["material_id"] = " NULL "
    result = ingest_csv_batch(csv_files(data))
    assert errors(result, "E-MISSING")

    data = rows()
    data["quality"][0]["date"] = data["quality"][0]["production_date"]
    fields = [*list(SCHEMAS["quality"]), "date"]
    result = ingest_csv_batch(csv_files(data, headers={"quality": fields}))
    assert any(issue.field == "production_date" for issue in errors(result, "E-SCHEMA"))

    files = csv_files(rows())
    files["inventory"] = ("../private.csv", files["inventory"][1])
    result = ingest_csv_batch(files)
    assert errors(result, "E-PACKAGE") and not result.accepted


def test_utf8_bom_is_logged_without_rejecting_valid_csv() -> None:
    files = csv_files(rows())
    filename, content = files["inventory"]
    files["inventory"] = (filename, b"\xef\xbb\xbf" + content)
    result = ingest_csv_batch(files)
    assert result.accepted
    assert len(errors(result, "N01")) == 1
    assert errors(result, "N01")[0].severity == "INFO"


def test_malformed_csv_and_blank_record_are_not_skipped() -> None:
    files = csv_files(rows())
    files["inventory"] = (
        "synthetic_inventory.csv",
        files["inventory"][1] + b"\n",
    )
    blank = ingest_csv_batch(files)
    assert any(issue.row == 3 for issue in errors(blank, "E-PARSE"))
    assert not blank.accepted

    files["inventory"] = (
        "synthetic_inventory.csv",
        files["inventory"][1] + b'"unterminated',
    )
    malformed = ingest_csv_batch(files)
    assert errors(malformed, "E-PARSE")
    assert malformed.total_observed_rows is None


def test_complete_synthetic_csv_batch_is_accepted() -> None:
    root = Path(__file__).resolve().parents[2] / "data" / "samples" / "clean"
    files = {
        dataset: (f"synthetic_{dataset}.csv", (root / f"synthetic_{dataset}.csv").read_bytes())
        for dataset in SCHEMAS
    }
    result = ingest_csv_batch(files)
    assert result.accepted
    assert result.total_observed_rows == 5220
    assert sum(len(records) for records in result.records.values()) == 5220
    assert result.error_count == 0
