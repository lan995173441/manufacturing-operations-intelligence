"""Decode CSV/XLSX into source-positioned candidates; business checks live in Domain."""

from __future__ import annotations

import csv
import hashlib
import io
import zipfile
from dataclasses import dataclass
from typing import Any
from uuid import uuid4
from xml.etree import ElementTree

from openpyxl import load_workbook

DATASETS = ("production_plan", "production_actual", "quality", "inventory")
MAX_BYTES = 10_000_000
MAX_ROWS = 10_000
MAX_EXPANDED_BYTES = 40_000_000
MAX_ARCHIVE_MEMBERS = 256
MAX_SHEET_CELLS = 200_000


def _preflight_xlsx(content: bytes) -> tuple[str, str] | None:
    """Bound decompression and worksheet complexity before openpyxl parses cells."""
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            members = archive.infolist()
            if (
                len(members) > MAX_ARCHIVE_MEMBERS
                or sum(member.file_size for member in members) > MAX_EXPANDED_BYTES
            ):
                return "E-LIMIT", "XLSX expanded content exceeds supported limits."
            expanded = 0
            cells = 0
            for member in members:
                if member.flag_bits & 1:
                    return "E-PARSE", "Encrypted XLSX files are unsupported."
                is_sheet = (
                    member.filename.startswith("xl/worksheets/")
                    and member.filename.endswith(".xml")
                )
                chunks = [] if is_sheet else None
                with archive.open(member) as source:
                    while chunk := source.read(1_000_000):
                        expanded += len(chunk)
                        if expanded > MAX_EXPANDED_BYTES:
                            return "E-LIMIT", "XLSX expanded content exceeds supported limits."
                        if chunks is not None:
                            chunks.append(chunk)
                if chunks is None:
                    continue
                try:
                    sheet_xml = b"".join(chunks)
                    if b"<!DOCTYPE" in sheet_xml or b"<!ENTITY" in sheet_xml:
                        return "E-PARSE", "XLSX worksheet DTDs are unsupported."
                    for _event, element in ElementTree.iterparse(
                        io.BytesIO(sheet_xml), events=("start",)
                    ):
                        tag = element.tag.rsplit("}", 1)[-1]
                        if tag == "c":
                            cells += 1
                            if cells > MAX_SHEET_CELLS:
                                return "E-LIMIT", "XLSX contains too many populated cells."
                        elif tag == "mergeCell":
                            return "E-SCHEMA", "Merged cells are unsupported."
                        element.clear()
                except ElementTree.ParseError:
                    return "E-PARSE", "Malformed XLSX worksheet XML."
    except (zipfile.BadZipFile, OSError, RuntimeError, ValueError):
        return "E-PARSE", "Cannot decode XLSX archive."
    return None


def _valid_filename(filename: str) -> bool:
    return (
        isinstance(filename, str)
        and 1 <= len(filename) <= 255
        and not any(character in filename for character in "/\\")
        and not any(ord(character) < 32 or ord(character) == 127 for character in filename)
    )


@dataclass(frozen=True)
class RawRow:
    row: int
    values: tuple[Any, ...]


@dataclass(frozen=True)
class ReadProblem:
    dataset: str | None
    row: int | None
    field: str | None
    code: str
    reason: str
    correction: str
    filename: str | None = None
    sheet: str | None = None


@dataclass(frozen=True)
class RawSource:
    dataset: str
    filename: str
    sheet: str | None
    source_id: str
    sha256: str
    kind: str
    headers: tuple[Any, ...]
    rows: tuple[RawRow, ...]
    total_rows: int | None
    problems: tuple[ReadProblem, ...] = ()
    had_bom: bool = False


def _source(
    dataset: str,
    filename: str,
    sheet: str | None,
    content: bytes,
    kind: str,
    headers: tuple[Any, ...],
    rows: tuple[RawRow, ...],
    total_rows: int | None,
    problems: tuple[ReadProblem, ...] = (),
    had_bom: bool = False,
) -> RawSource:
    return RawSource(
        dataset,
        filename,
        sheet,
        str(uuid4()),
        hashlib.sha256(content).hexdigest(),
        kind,
        headers,
        rows,
        total_rows,
        problems,
        had_bom,
    )


def _problem(
    dataset: str | None,
    filename: str,
    code: str,
    reason: str,
    *,
    row: int | None = None,
    field: str | None = None,
    sheet: str | None = None,
) -> ReadProblem:
    return ReadProblem(
        dataset, row, field, code, reason, "Correct the source file and retry.", filename, sheet
    )


def read_csv(dataset: str, filename: str, content: bytes) -> RawSource:
    """Read an explicitly assigned UTF-8 CSV without coercing any value."""
    if not _valid_filename(filename):
        return _source(
            dataset,
            filename,
            None,
            content,
            "csv",
            (),
            (),
            None,
            (_problem(dataset, filename, "E-PACKAGE", "Filename must be a safe basename."),),
        )
    if not filename.lower().endswith(".csv"):
        return _source(
            dataset,
            filename,
            None,
            content,
            "csv",
            (),
            (),
            None,
            (_problem(dataset, filename, "E-PACKAGE", "Expected a .csv file."),),
        )
    if len(content) > MAX_BYTES:
        return _source(
            dataset,
            filename,
            None,
            content,
            "csv",
            (),
            (),
            None,
            (_problem(dataset, filename, "E-LIMIT", "File exceeds 10,000,000 bytes."),),
        )
    try:
        decoded = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        return _source(
            dataset,
            filename,
            None,
            content,
            "csv",
            (),
            (),
            None,
            (_problem(dataset, filename, "E-PARSE", "CSV is not valid UTF-8."),),
        )

    headers: tuple[str, ...] = ()
    rows: list[RawRow] = []
    problems: list[ReadProblem] = []
    reader = csv.reader(io.StringIO(decoded, newline=""), strict=True)
    try:
        headers = tuple(next(reader))
        if reader.line_num != 1:
            problems.append(
                _problem(dataset, filename, "E-PARSE", "Multiline header is unsupported.")
            )
        previous_line = reader.line_num
        for values in reader:
            source_row = previous_line + 1
            if reader.line_num != source_row:
                problems.append(
                    _problem(
                        dataset,
                        filename,
                        "E-PARSE",
                        "Embedded line break is unsupported.",
                        row=source_row,
                    )
                )
                break
            previous_line = reader.line_num
            if len(values) != len(headers):
                problems.append(
                    _problem(
                        dataset,
                        filename,
                        "E-PARSE",
                        "CSV field count differs from header.",
                        row=source_row,
                    )
                )
            rows.append(RawRow(source_row, tuple(values)))
            if len(rows) > MAX_ROWS:
                problems.append(
                    _problem(dataset, filename, "E-LIMIT", "More than 10,000 data rows.")
                )
                break
    except StopIteration:
        problems.append(_problem(dataset, filename, "E-SCHEMA", "CSV has no header row."))
    except csv.Error as exc:
        problems.append(
            _problem(
                dataset,
                filename,
                "E-PARSE",
                f"Malformed CSV near physical line {reader.line_num}: {exc}",
                row=reader.line_num or None,
            )
        )
    return _source(
        dataset,
        filename,
        None,
        content,
        "csv",
        headers,
        tuple(rows),
        None if any(p.code in {"E-PARSE", "E-LIMIT"} for p in problems) else len(rows),
        tuple(problems),
        content.startswith(b"\xef\xbb\xbf"),
    )


def read_workbook(
    filename: str, content: bytes
) -> tuple[tuple[RawSource, ...], tuple[ReadProblem, ...]]:
    """Read exactly four canonical worksheets, retaining cell types and physical rows."""
    if not _valid_filename(filename):
        return (), (_problem(None, filename, "E-PACKAGE", "Filename must be a safe basename."),)
    if not filename.lower().endswith(".xlsx"):
        return (), (_problem(None, filename, "E-PACKAGE", "Expected a .xlsx workbook."),)
    if len(content) > MAX_BYTES:
        return (), (_problem(None, filename, "E-LIMIT", "File exceeds 10,000,000 bytes."),)
    preflight = _preflight_xlsx(content)
    if preflight is not None:
        code, reason = preflight
        return (), (_problem(None, filename, code, reason),)
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except Exception as exc:
        return (), (
            _problem(None, filename, "E-PARSE", f"Cannot decode XLSX: {type(exc).__name__}."),
        )

    sources: list[RawSource] = []
    problems: list[ReadProblem] = []
    try:
        if set(workbook.sheetnames) != set(DATASETS) or len(workbook.sheetnames) != 4:
            problems.append(
                _problem(
                    None,
                    filename,
                    "E-PACKAGE",
                    "Workbook must contain exactly the four canonical sheets.",
                )
            )
        for dataset in DATASETS:
            if dataset not in workbook:
                continue
            sheet = workbook[dataset]
            local: list[ReadProblem] = []
            if (sheet.max_row is not None and sheet.max_row > MAX_ROWS + 1) or (
                sheet.max_column is not None and sheet.max_column > 100
            ):
                local.append(
                    _problem(
                        dataset,
                        filename,
                        "E-LIMIT",
                        "Worksheet dimensions exceed the supported 10,000 rows or 100 columns.",
                        sheet=dataset,
                    )
                )
                sources.append(
                    _source(dataset, filename, dataset, content, "xlsx", (), (), None, tuple(local))
                )
                continue
            iterator = sheet.iter_rows()
            header_cells = next(iterator, ())
            headers = tuple(cell.value for cell in header_cells)
            rows: list[RawRow] = []
            for row_num, cells in enumerate(iterator, start=2):
                if row_num > MAX_ROWS + 1 or len(cells) > 100:
                    local.append(
                        _problem(
                            dataset, filename, "E-LIMIT",
                            "Worksheet dimensions exceed the supported 10,000 rows or 100 columns.",
                            sheet=dataset,
                        )
                    )
                    break
                for cell in cells:
                    if cell.data_type == "f":
                        local.append(
                            _problem(
                                dataset,
                                filename,
                                "E-PARSE",
                                "Formula cells are unsupported.",
                                row=row_num,
                                field=headers[cell.column - 1]
                                if cell.column <= len(headers) else None,
                                sheet=dataset,
                            )
                        )
                rows.append(RawRow(row_num, tuple(cell.value for cell in cells)))
                if len(rows) > MAX_ROWS:
                    local.append(
                        _problem(
                            dataset,
                            filename,
                            "E-LIMIT",
                            "More than 10,000 data rows.",
                            sheet=dataset,
                        )
                    )
                    break
            while rows and all(value is None for value in rows[-1].values):
                rows.pop()
            sources.append(
                _source(
                    dataset,
                    filename,
                    dataset,
                    content,
                    "xlsx",
                    headers,
                    tuple(rows),
                    None if len(rows) > MAX_ROWS else len(rows),
                    tuple(local),
                )
            )
    except Exception as exc:
        return (), (
            _problem(None, filename, "E-PARSE", f"Cannot read XLSX cells: {type(exc).__name__}."),
        )
    finally:
        workbook.close()
    return tuple(sources), tuple(problems)
