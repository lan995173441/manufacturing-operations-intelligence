"""Present service-calculated manufacturing results as Excel and PDF reports."""

from __future__ import annotations

from collections.abc import Mapping
from dataclasses import dataclass
from decimal import ROUND_HALF_UP, Decimal, localcontext
from fractions import Fraction
from io import BytesIO
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.graphics.shapes import Drawing, Line, PolyLine, String
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from manufacturing_operations_intelligence.data.repository import BatchIdentity
from manufacturing_operations_intelligence.domain.analytics import (
    DashboardSeries,
    KpiResult,
    KpiScope,
)
from manufacturing_operations_intelligence.domain.anomalies import DetectionResult

LABELS = {
    "KPI-PA": "Production attainment",
    "KPI-GY": "Good yield",
    "KPI-SR": "Scrap rate",
    "KPI-TD": "Total downtime",
    "KPI-CO": "Completed orders",
    "KPI-SA": "Schedule adherence",
    "KPI-IR": "Materials below safety stock",
    "KPI-OL": "Output by production line",
}
KPI_DEFINITION_SUMMARIES = (
    ("Production attainment", "Gross actual output divided by planned output."),
    ("Good yield", "Final good quantity divided by actual output."),
    ("Scrap rate", "Final scrap quantity divided by actual output."),
    ("Total downtime", "Sum of recorded downtime in line-minutes."),
    ("Completed orders", "Unavailable: the contract has no completion events or order census."),
    ("Schedule adherence", "Unavailable: the contract has no due or completion evidence."),
    (
        "Materials below safety stock",
        "Count of materials strictly below safety stock as of end date.",
    ),
    ("Output by production line", "Actual output grouped by production line."),
)
INK = "17324D"


@dataclass(frozen=True)
class ReportPayload:
    identity: BatchIdentity
    scope: KpiScope
    metrics: Mapping[str, KpiResult]
    series: DashboardSeries
    detection: DetectionResult | None
    generated_utc: str


@dataclass(frozen=True)
class ReportArtifacts:
    identity: BatchIdentity
    scope: KpiScope
    excel: bytes | None
    pdf: bytes | None
    excel_error: str | None = None
    pdf_error: str | None = None


@dataclass(frozen=True)
class ReportNarrative:
    """Display-only narrative supplied by the application layer after identity checks."""

    source: str
    observations: tuple[str, ...]
    management_attention: tuple[str, ...]
    limitations: tuple[str, ...]


def _verify(payload: ReportPayload) -> None:
    if set(payload.metrics) != set(LABELS):
        raise ValueError("Report requires all eight calculated KPI results.")
    definition_version = payload.metrics["KPI-PA"].definition_version
    if any(
        metric.kpi_id != key
        or metric.batch_id != payload.identity.batch_id
        or metric.fingerprint != payload.identity.fingerprint
        or metric.contract_version != payload.identity.contract_version
        or metric.definition_version != definition_version
        or metric.scope != payload.scope
        for key, metric in payload.metrics.items()
    ):
        raise ValueError("KPI result does not match report identity, definition or filters.")
    if payload.detection is not None and (
        payload.detection.batch_id != payload.identity.batch_id
        or payload.detection.scope != payload.scope
    ):
        raise ValueError("Anomaly result does not match report batch or filters.")


def _value(metric: KpiResult) -> str:
    return metric.display_value if metric.status == "valid" and metric.display_value else "N/A"


def _scope(payload: ReportPayload) -> list[tuple[str, str]]:
    selection = payload.scope
    return [
        ("Report period", f"{selection.start_date} to {selection.end_date} (inclusive)"),
        ("Production lines", ", ".join(selection.line_ids) if selection.line_ids else "All"),
        ("Finished products", ", ".join(selection.product_ids) if selection.product_ids else "All"),
        ("Production shifts", ", ".join(selection.shift_ids) if selection.shift_ids else "All"),
        ("Batch ID", payload.identity.batch_id),
        ("Imported UTC", payload.identity.imported_at_utc),
        ("Generated UTC", payload.generated_utc),
        ("Data contract", payload.identity.contract_version),
        ("KPI definitions", payload.metrics["KPI-PA"].definition_version),
        (
            "Inventory scope",
            "Latest eligible material snapshots as of the end date; "
            "production line and product filters do not apply.",
        ),
    ]


def _pdf_scope(payload: ReportPayload) -> list[tuple[str, str]]:
    """Group short identifiers, while bounding every splittable PDF row."""
    rows = _scope(payload)
    for label, identifiers in (
        ("Production lines", payload.scope.line_ids),
        ("Finished products", payload.scope.product_ids),
        ("Production shifts", payload.scope.shift_ids),
    ):
        if identifiers:
            index = next(index for index, row in enumerate(rows) if row[0] == label)
            groups: list[str] = []
            group: list[str] = []
            for identifier in identifiers:
                proposed = ", ".join((*group, identifier))
                if group and (len(proposed) > 90 or len(group) == 8):
                    groups.append(", ".join(group))
                    group = []
                group.append(identifier)
            if group:
                groups.append(", ".join(group))
            rows[index:index + 1] = [
                (label if position == 0 else "", text)
                for position, text in enumerate(groups)
            ]
    return rows


def _alert_state(payload: ReportPayload) -> str:
    result = payload.detection
    if result is None:
        return "Not evaluated: configure a downtime threshold."
    if result.status in {"invalid_data", "unsupported_schema"}:
        return f"Unavailable: {result.reason or result.status}"
    if result.status == "no_data" or result.evaluated_count == 0:
        return "No checks completed for this scope."
    if result.status == "partial_coverage" or result.skipped_count:
        return "Partial coverage: some checks were skipped."
    return "Completed"


def _summary(payload: ReportPayload) -> list[str]:
    metrics = payload.metrics
    lines = [
        f"For {payload.scope.start_date} to {payload.scope.end_date}, production attainment "
        f"is {_value(metrics['KPI-PA'])}, good yield is {_value(metrics['KPI-GY'])}, "
        f"and scrap rate is {_value(metrics['KPI-SR'])}.",
        f"Recorded downtime is {_value(metrics['KPI-TD'])}. "
        f"Materials below safety stock: {_value(metrics['KPI-IR'])}.",
    ]
    if payload.detection is None:
        lines.append("Anomalies were not evaluated: a downtime threshold is not configured.")
    else:
        result = payload.detection
        lines.append(
            f"Detected anomalies: {result.anomaly_count}; completed checks: "
            f"{result.evaluated_count}; skipped checks: {result.skipped_count}. "
            f"{_alert_state(payload)}"
        )
    lines.append(
        "Completed orders and schedule adherence are unavailable under the current data "
        "contract; this report makes no completion or on-time claim."
    )
    return lines


def _narrative_sections(
    payload: ReportPayload, narrative: ReportNarrative | None
) -> tuple[str, tuple[tuple[str, tuple[str, ...]], ...]]:
    """Keep the default deterministic narrative while allowing a verified current draft."""
    if narrative is None:
        return (
            "Deterministic summary from calculated results; no AI-generated claims.",
            (("Observations", tuple(_summary(payload))),),
        )
    return (
        narrative.source,
        (
            ("Observations", narrative.observations),
            ("Management attention", narrative.management_attention),
            ("Limitations", narrative.limitations),
        ),
    )


def _anomalies(payload: ReportPayload) -> list[tuple[object, ...]]:
    if payload.detection is None:
        return []
    return [
        (
            item.date_start,
            item.date_end,
            item.severity,
            item.type.replace("_", " ").title(),
            item.entity.key,
            _number_text(item.observed_value),
            item.comparator,
            _number_text(item.threshold),
            item.unit,
            item.explanation,
        )
        for item in sorted(
            payload.detection.anomalies,
            key=lambda item: (item.date_start, item.type, item.entity.key),
        )
    ]


def _number_text(value: Fraction | Decimal) -> str:
    """Round exact anomaly evidence for display, without changing rule outcomes."""
    if isinstance(value, Fraction):
        with localcontext() as context:
            context.prec = 40
            decimal = Decimal(value.numerator) / Decimal(value.denominator)
        return str(decimal.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))
    return str(value)


def _append(sheet, values: tuple[object, ...]) -> None:
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            # Explicit literal type prevents an uploaded identifier from becoming a formula.
            cell.data_type = "s"


def _sheet(book: Workbook, name: str, heading: tuple[str, ...]):
    sheet = book.create_sheet(name)
    _append(sheet, (name + " | Manufacturing Operations Intelligence",))
    sheet["A1"].font = Font(size=15, bold=True, color=INK)
    _append(sheet, heading)
    for cell in sheet[2]:
        cell.fill = PatternFill("solid", fgColor=INK)
        cell.font = Font(bold=True, color="FFFFFF")
    sheet.freeze_panes = "A3"
    sheet.sheet_view.showGridLines = False
    return sheet


def _finish(sheet) -> None:
    for column in sheet.columns:
        cells = list(column)
        width = min(56, max(14, *(len(str(cell.value or "")) + 2 for cell in cells)))
        sheet.column_dimensions[get_column_letter(cells[0].column)].width = width
    for row in sheet.iter_rows(min_row=3):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def _excel_chart(sheet, kind: str, count: int, anchor: str) -> None:
    if not count:
        return
    chart = LineChart() if kind == "line" else BarChart()
    chart.title = "Planned vs actual output" if kind == "line" else "Good vs scrap output"
    chart.y_axis.title = "Finished pieces (ea)"
    chart.add_data(
        Reference(sheet, min_col=2, max_col=3, min_row=2, max_row=count + 2),
        titles_from_data=True,
    )
    chart.set_categories(Reference(sheet, min_col=1, min_row=3, max_row=count + 2))
    chart.width, chart.height = 21, 10
    sheet.add_chart(chart, anchor)


def render_excel(payload: ReportPayload, narrative: ReportNarrative | None = None) -> bytes:
    """Export final values, prepared series and complete detail without KPI formulas."""
    _verify(payload)
    book = Workbook()
    book.remove(book.active)
    overview = _sheet(book, "Overview", ("Item", "Value", "Status / explanation"))
    for label, value in _scope(payload):
        _append(overview, (label, value, ""))
    _append(overview, ("Anomaly checks", _alert_state(payload), ""))
    for key, label in LABELS.items():
        metric = payload.metrics[key]
        _append(overview, (label, _value(metric), metric.reason or metric.status))

    production = _sheet(
        book, "Production", ("Date", "Planned (ea)", "Actual (ea)", "Downtime (line-minutes)")
    )
    quality = _sheet(book, "Quality", ("Date", "Good (ea)", "Scrap (ea)"))
    for point in payload.series.daily:
        _append(
            production,
            (point.group, point.planned_qty, point.actual_qty, float(point.downtime_minutes)),
        )
        _append(quality, (point.group, point.good_qty, point.scrap_qty))
    if not payload.series.daily:
        _append(production, ("No selected production records",))
        _append(quality, ("No selected quality records",))
    _excel_chart(production, "line", len(payload.series.daily), "F3")
    _excel_chart(quality, "bar", len(payload.series.daily), "E3")
    _append(production, ("Production line", "Planned (ea)", "Actual (ea)"))
    for point in payload.series.by_line:
        _append(production, (point.group, point.planned_qty, point.actual_qty))

    inventory = _sheet(
        book,
        "Inventory",
        (
            "Material",
            "On hand",
            "Safety stock",
            "Unit",
            "Observed on",
            "Carried forward",
            "Below safety stock",
            "Coverage",
        ),
    )
    for item in payload.metrics["KPI-IR"].details:
        _append(
            inventory,
            (
                item["material_id"],
                item.get("inventory_qty"),
                item.get("safety_stock"),
                item.get("qty_unit"),
                item.get("snapshot_date"),
                item.get("carried_forward"),
                item.get("below_safety_stock"),
                item["status"],
            ),
        )
    if not payload.metrics["KPI-IR"].details:
        _append(inventory, ("No eligible material observations",))

    anomaly_sheet = _sheet(
        book,
        "Anomalies",
        (
            "From",
            "Through",
            "Severity",
            "Rule",
            "Entity",
            "Observed",
            "Comparison",
            "Threshold",
            "Unit",
            "Explanation",
        ),
    )
    anomaly_rows = _anomalies(payload)
    for row in anomaly_rows:
        _append(anomaly_sheet, row)
    if not anomaly_rows:
        _append(anomaly_sheet, (_alert_state(payload),))

    audit = _sheet(book, "Data Quality", ("Measure", "Status", "Reason", "Coverage evidence"))
    for key, label in LABELS.items():
        metric = payload.metrics[key]
        coverage = "; ".join(
            f"{name.replace('_', ' ')}: {value}" for name, value in metric.coverage.items()
        )
        _append(audit, (label, metric.status, metric.reason or "", coverage))
    _append(audit, ("Anomaly checks", _alert_state(payload), "", ""))
    if payload.detection is not None:
        _append(audit, ("Completed checks", payload.detection.evaluated_count, "", ""))
        _append(audit, ("Skipped checks", payload.detection.skipped_count, "", ""))
        for skipped in payload.detection.skipped:
            _append(audit, (skipped.type, "skipped", skipped.reason, skipped.entity.key))

    summary = _sheet(book, "Summary", ("Section", "Management summary"))
    source, sections = _narrative_sections(payload, narrative)
    _append(summary, ("Summary type", source))
    for heading, lines in sections:
        for line in lines:
            _append(summary, (heading, line))
    for sheet in book:
        _finish(sheet)
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def _p(value: object, style: ParagraphStyle) -> Paragraph:
    return Paragraph(escape(str(value if value is not None else "")), style)


def _table(rows: list[tuple[object, ...]], widths: list[int], styles) -> Table:
    cells = [
        [_p(value, styles["HeaderCell"] if index == 0 else styles["Cell"]) for value in row]
        for index, row in enumerate(rows)
    ]
    table = Table(cells, colWidths=widths, repeatRows=1, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#17324D")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F6F8")]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    return table


def _plot(points, title: str, fields: tuple[tuple[str, str, str], ...]) -> Drawing:
    """Scale already prepared daily values into vector coordinates; no KPI math."""
    drawing = Drawing(500, 185)
    drawing.add(String(8, 169, title, fontName="Helvetica-Bold", fontSize=10))
    x0, y0, width, height = 48, 35, 427, 112
    drawing.add(Line(x0, y0, x0 + width, y0, strokeColor=colors.grey))
    drawing.add(Line(x0, y0, x0, y0 + height, strokeColor=colors.grey))
    if not points:
        drawing.add(String(170, 86, "No selected records", fontSize=10))
        return drawing
    top = max(1, *(float(getattr(point, field)) for point in points for _, field, _ in fields))
    drawing.add(String(3, y0 + height, f"{top:,.0f}", fontSize=7))
    drawing.add(String(x0, 20, points[0].group, fontSize=7))
    drawing.add(String(x0 + width - 55, 20, points[-1].group, fontSize=7))
    for index, (label, field, color) in enumerate(fields):
        coordinates = []
        for position, point in enumerate(points):
            coordinates.extend(
                (
                    x0 + width * position / max(1, len(points) - 1),
                    y0 + height * float(getattr(point, field)) / top,
                )
            )
        drawing.add(PolyLine(coordinates, strokeColor=colors.HexColor(color), strokeWidth=1.5))
        drawing.add(
            String(58 + index * 160, 155, label, fontSize=8, fillColor=colors.HexColor(color))
        )
    return drawing


def render_pdf(payload: ReportPayload, narrative: ReportNarrative | None = None) -> bytes:
    """Generate a paginated review PDF from the same immutable report payload."""
    _verify(payload)
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="ReportTitle",
            parent=styles["Title"],
            fontSize=18,
            leading=23,
            textColor=colors.HexColor("#17324D"),
        )
    )
    styles.add(
        ParagraphStyle(
            name="Section",
            parent=styles["Heading2"],
            fontSize=12,
            leading=16,
            textColor=colors.HexColor("#17324D"),
            spaceBefore=14,
        )
    )
    styles.add(
        ParagraphStyle(
            name="Cell", parent=styles["Normal"], fontSize=8, leading=10, splitLongWords=1
        )
    )
    styles.add(
        ParagraphStyle(
            name="HeaderCell",
            parent=styles["Cell"],
            fontName="Helvetica-Bold",
            textColor=colors.white,
        )
    )
    story = [
        Paragraph("Manufacturing Operations Intelligence", styles["ReportTitle"]),
        Paragraph("Management report", styles["Section"]),
        _table([("Scope", "Value"), *_pdf_scope(payload)], [130, 370], styles),
        Paragraph("Executive KPI summary", styles["Section"]),
        _table(
            [("Indicator", "Result", "Status / reason")]
            + [
                (
                    LABELS[key],
                    _value(payload.metrics[key]),
                    payload.metrics[key].reason or payload.metrics[key].status,
                )
                for key in LABELS
                if key != "KPI-OL"
            ],
            [155, 120, 225],
            styles,
        ),
        Paragraph("KPI definitions", styles["Section"]),
        _table(
            [("Indicator", "Current candidate definition"), *KPI_DEFINITION_SUMMARIES],
            [155, 345],
            styles,
        ),
        KeepTogether(
            [
                Paragraph("Production analysis", styles["Section"]),
                _plot(
                    payload.series.daily,
                    "Planned vs actual output by day (ea)",
                    (("Planned", "planned_qty", "#2563EB"), ("Actual", "actual_qty", "#0F766E")),
                ),
            ]
        ),
        KeepTogether(
            [
                Paragraph("Quality analysis", styles["Section"]),
                _plot(
                    payload.series.daily,
                    "Final good vs scrap output by day (ea)",
                    (("Good", "good_qty", "#0F766E"), ("Scrap", "scrap_qty", "#D97706")),
                ),
            ]
        ),
        Paragraph("Inventory risks", styles["Section"]),
        _p(
            "Materials below safety stock: " + _value(payload.metrics["KPI-IR"]) + ". "
            "Inventory uses the latest eligible snapshot at or before the end date.",
            styles["Normal"],
        ),
    ]
    risks = [
        item
        for item in payload.metrics["KPI-IR"].details
        if item.get("below_safety_stock") or item.get("status") != "observed"
    ]
    if risks:
        story.append(
            _table(
                [("Material", "On hand", "Safety stock", "Unit", "Observed / coverage")]
                + [
                    (
                        item["material_id"],
                        item.get("inventory_qty")
                        if item.get("inventory_qty") is not None
                        else "N/A",
                        item.get("safety_stock") if item.get("safety_stock") is not None else "N/A",
                        item.get("qty_unit", ""),
                        item.get("snapshot_date") or item.get("status", ""),
                    )
                    for item in risks[:20]
                ],
                [110, 85, 95, 55, 155],
                styles,
            )
        )
        if len(risks) > 20:
            story.append(
                _p(
                    f"Showing 20 of {len(risks)} inventory risks. See Excel for all.",
                    styles["Normal"],
                )
            )
    else:
        story.append(
            _p("No inventory items are flagged in the available observations.", styles["Normal"])
        )

    story.extend(
        [
            Paragraph("Detected anomalies", styles["Section"]),
            _p(_alert_state(payload), styles["Normal"]),
        ]
    )
    if payload.detection is not None:
        result = payload.detection
        story.append(
            _p(
                f"Detected: {result.anomaly_count}; completed: {result.evaluated_count}; "
                f"skipped: {result.skipped_count}.",
                styles["Normal"],
            )
        )
    anomaly_rows = _anomalies(payload)
    if anomaly_rows:
        story.append(
            _table(
                [("Date", "Severity", "Rule", "Entity", "Observed / threshold")]
                + [
                    (row[0], row[2], row[3], row[4], f"{row[5]} {row[6]} {row[7]} {row[8]}")
                    for row in anomaly_rows[:20]
                ],
                [68, 60, 110, 140, 122],
                styles,
            )
        )
        if len(anomaly_rows) > 20:
            story.append(
                _p(
                    f"Showing 20 of {len(anomaly_rows)} anomalies. See Excel for all.",
                    styles["Normal"],
                )
            )
    elif payload.detection is not None:
        story.append(_p("No anomalies detected in available checks.", styles["Normal"]))
    story.append(Paragraph("Data quality and coverage", styles["Section"]))
    for key, label in LABELS.items():
        metric = payload.metrics[key]
        if metric.status != "valid":
            story.append(_p(f"{label}: {metric.status}. {metric.reason or ''}", styles["Normal"]))
    if payload.detection is not None and payload.detection.skipped_count:
        story.append(
            _p(
                f"{payload.detection.skipped_count} checks were skipped. "
                "See Excel Data Quality for reasons.",
                styles["Normal"],
            )
        )
    story.append(Paragraph("Management summary", styles["Section"]))
    source, sections = _narrative_sections(payload, narrative)
    story.append(_p(source, styles["Normal"]))
    for heading, lines in sections:
        story.append(_p(f"{heading}:", styles["Normal"]))
        for line in lines:
            story.extend((_p(line, styles["Normal"]), Spacer(1, 5)))

    def footer(canvas, document) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.drawString(
            54, 33, f"Batch {payload.identity.batch_id[:12]} | Generated {payload.generated_utc}"
        )
        canvas.drawRightString(555, 33, f"Page {document.page}")
        canvas.restoreState()

    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=(612, 792),
        leftMargin=54,
        rightMargin=58,
        topMargin=46,
        bottomMargin=54,
        pageCompression=0,
    )
    document.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()
